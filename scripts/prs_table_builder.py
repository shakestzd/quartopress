"""Reusable PRS three-line table builder for Word documents.

Generates publication-ready tables with PRS journal formatting:
- Thick top and bottom rules
- Thin rule under header row
- No internal row borders, no vertical lines, no shading

Usage:
    from prs_table_builder import TableSpec, build_prs_document

    # From explicit rows
    spec = TableSpec(
        label="Table 1. Patient Demographics",
        headers=["Characteristic", "Value"],
        rows=[("Age (years)", "45.2 +/- 12.3"), ...],
    )

    # From a DataFrame
    spec = TableSpec.from_dataframe(
        df, label="Table 2. Performance Metrics", col1_bold=False,
    )

    # Build single table or full document
    build_prs_document(
        title_bold="Table 1. ",
        title_text="Description of the table.",
        tables=[spec],
        output_path="Table_1.docx",
    )
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt


# =============================================================================
# Low-Level Formatting
# =============================================================================


def _set_cell_border(cell, **kwargs):
    """Set border properties on a single cell edge.

    Args:
        cell: python-docx table cell.
        **kwargs: Edge name -> {"sz": int, "val": str, "color": str}.
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}></w:tcBorders>')
    for edge, attrs in kwargs.items():
        element = parse_xml(
            f'<w:{edge} {nsdecls("w")} w:val="{attrs["val"]}" '
            f'w:sz="{attrs["sz"]}" w:space="0" w:color="{attrs["color"]}"/>'
        )
        tcBorders.append(element)
    tcPr.append(tcBorders)


def _apply_prs_table_borders(table):
    """Apply PRS three-line borders to a table object."""
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}></w:tblPr>')
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="000000"/>'
        '  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="000000"/>'
        '  <w:left w:val="none" w:sz="0" w:space="0" w:color="000000"/>'
        '  <w:right w:val="none" w:sz="0" w:space="0" w:color="000000"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)


def _apply_header_bottom_border(row):
    """Add thin bottom border under header row."""
    border = {"sz": 8, "val": "single", "color": "000000"}
    for cell in row.cells:
        _set_cell_border(cell, bottom=border)


def _set_cell_text(cell, text: str, bold=False, italic=False, size=9):
    """Write formatted text into a table cell."""
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


# =============================================================================
# TableSpec
# =============================================================================


@dataclass
class TableSpec:
    """Declarative specification for a PRS-style table.

    Attributes:
        label: Table title shown above the table (e.g. "Table 1. Demographics").
        headers: Column header texts.
        rows: List of tuples, one per data row. Length must match headers.
        col_widths: Column widths in inches. Defaults to equal distribution.
        col1_bold: Whether to bold the first column in data rows.
        section_headers: Optional dict mapping row index -> section header text,
            inserted BEFORE that row (for multi-section tables).
    """

    label: str
    headers: list[str]
    rows: list[tuple]
    col_widths: tuple[float, ...] | None = None
    col1_bold: bool = True
    section_headers: dict[int, str] = field(default_factory=dict)

    @classmethod
    def from_dataframe(
        cls,
        df: pd.DataFrame,
        label: str,
        col1_bold: bool = True,
        col_widths: tuple[float, ...] | None = None,
    ) -> TableSpec:
        """Create a TableSpec from a pandas DataFrame.

        Args:
            df: Source DataFrame. Column names become headers.
            label: Table title.
            col1_bold: Bold first column values.
            col_widths: Optional column widths in inches.
        """
        headers = list(df.columns)
        rows = [tuple(row) for row in df.itertuples(index=False, name=None)]
        return cls(
            label=label,
            headers=headers,
            rows=rows,
            col1_bold=col1_bold,
            col_widths=col_widths,
        )

    @classmethod
    def from_csv(
        cls,
        csv_path: str | Path,
        label: str,
        col1_bold: bool = True,
        col_widths: tuple[float, ...] | None = None,
    ) -> TableSpec:
        """Create a TableSpec directly from a CSV file.

        Args:
            csv_path: Path to CSV file.
            label: Table title.
            col1_bold: Bold first column values.
            col_widths: Optional column widths in inches.
        """
        df = pd.read_csv(csv_path)
        return cls.from_dataframe(df, label=label, col1_bold=col1_bold, col_widths=col_widths)


# =============================================================================
# Table Builder
# =============================================================================


def build_prs_table(
    doc: Document,
    spec: TableSpec,
    header_size: int = 10,
    body_size: int = 9,
) -> None:
    """Build a single PRS three-line table in a Document from a TableSpec.

    Args:
        doc: python-docx Document to add the table to.
        spec: Table specification.
        header_size: Font size for header row (pt).
        body_size: Font size for body rows (pt).
    """
    ncols = len(spec.headers)

    # Default column widths: distribute 6.5 inches evenly
    col_widths = spec.col_widths or tuple(6.5 / ncols for _ in range(ncols))

    # Label
    p = doc.add_paragraph()
    p.add_run(spec.label).bold = True

    # Create table
    table = doc.add_table(rows=1, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    # Borders
    _apply_prs_table_borders(table)

    # Header row
    for i, text in enumerate(spec.headers):
        _set_cell_text(table.rows[0].cells[i], text, bold=True, size=header_size)
    _apply_header_bottom_border(table.rows[0])

    # Data rows
    for idx, row_data in enumerate(spec.rows):
        # Insert section header if defined
        if idx in spec.section_headers:
            sec_row = table.add_row()
            if ncols > 1:
                sec_row.cells[0].merge(sec_row.cells[ncols - 1])
            _set_cell_text(sec_row.cells[0], spec.section_headers[idx], bold=True, size=header_size)

        row = table.add_row()
        for col_idx, value in enumerate(row_data):
            bold = spec.col1_bold and col_idx == 0
            _set_cell_text(row.cells[col_idx], str(value), bold=bold, size=body_size)

    # Spacer
    doc.add_paragraph()


def build_prs_document(
    title_bold: str,
    title_text: str,
    tables: list[TableSpec],
    output_path: str | Path,
    font_name: str = "Times New Roman",
    font_size: int = 11,
    landscape: bool = False,
) -> Path:
    """Build a complete PRS-compliant Word document from table specs.

    Args:
        title_bold: Bold portion of the document title.
        title_text: Normal portion of the document title.
        tables: List of TableSpec objects to render.
        output_path: Output .docx file path.
        font_name: Document font.
        font_size: Document font size (pt).
        landscape: If True, set page orientation to landscape.

    Returns:
        Path to the saved document.
    """
    doc = Document()

    # Page orientation
    if landscape:
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        # Swap width and height for landscape
        section.page_width, section.page_height = section.page_height, section.page_width

    style = doc.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(font_size)

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title_para.add_run(title_bold)
    run.bold = True
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run = title_para.add_run(title_text)
    run.font.name = font_name
    run.font.size = Pt(font_size)

    doc.add_paragraph()

    for spec in tables:
        build_prs_table(doc, spec)

    output_path = Path(output_path)
    doc.save(output_path)
    print(f"Saved: {output_path}")
    return output_path
